#!/usr/bin/env python3
"""Guard duty scheduler.

Usage:
    python guard_duty.py config.json [--out DIR]

Exit codes: 0 = schedule written, 2 = not enough people for a night window
(nothing written; read the SHORTAGE report and ask the user how to fix it),
3 = some slots could not be filled (files written, holes marked "???").
"""
import argparse
import datetime as dt
import json
import os
import random
import sys
import time
from collections import defaultdict

SG, YUS, DAR, PAT = "ש.ג.", "יוסטור", "דרומית", "תצפית/פטרול"
STANDS = [SG, YUS, DAR, PAT]
ALIASES = {
    "sg": SG, "ש.ג": SG, "שג": SG, "ש\"ג": SG, "gate": SG,
    "yustor": YUS, "יוסטור": YUS,
    "dromit": DAR, "south": DAR, "דרומית": DAR,
    "patrol": PAT, "פטרול": PAT, "תצפית": PAT, "תצפית/פטרול": PAT,
}
DAY_START_HOUR = 7          # a "day" on the sheet runs 07:00 -> 07:00
NIGHT_START, NIGHT_END = 23, 7
H = dt.timedelta(hours=1)
HEB_DAYS = ["שני", "שלישי", "רביעי", "חמישי", "שישי", "שבת", "ראשון"]


def stand(name):
    if name in STANDS:
        return name
    key = str(name).strip().lower()
    if key in ALIASES:
        return ALIASES[key]
    raise SystemExit(f"Unknown stand: {name!r}. Use one of {STANDS}")


def parse_dt(s):
    return dt.datetime.strptime(s.strip(), "%Y-%m-%d %H:%M")


def night_id(t):
    """The date whose 23:00 opens the night window containing t, else None."""
    if t.hour >= NIGHT_START:
        return t.date()
    if t.hour < NIGHT_END:
        return t.date() - dt.timedelta(days=1)
    return None


def day_id(t):
    return (t - DAY_START_HOUR * H).date()


class Slot:
    __slots__ = ("idx", "stand", "start", "end", "night", "cancelled")

    def __init__(self, idx, st, start, end):
        self.idx, self.stand, self.start, self.end = idx, st, start, end
        self.night = night_id(start)
        self.cancelled = False

    @property
    def hours(self):
        return (self.end - self.start) / H


def build_slots(start, end):
    raw = []
    d0 = dt.datetime.combine(day_id(start), dt.time(DAY_START_HOUR))
    while d0 < end:
        for st in (SG, YUS, DAR):
            for k in range(12):
                raw.append((st, d0 + 2 * k * H, d0 + 2 * (k + 1) * H))
        for k in range(4):  # 07-11, 11-15, 15-19, 19-23
            raw.append((PAT, d0 + 4 * k * H, d0 + 4 * (k + 1) * H))
        for k in range(4):  # 23-01, 01-03, 03-05, 05-07
            s = d0 + 16 * H + 2 * k * H
            raw.append((PAT, s, s + 2 * H))
        d0 += 24 * H
    slots = []
    for st, s, e in raw:
        s, e = max(s, start), min(e, end)
        if e > s:
            slots.append(Slot(len(slots), st, s, e))
    slots.sort(key=lambda x: (x.start, STANDS.index(x.stand)))
    for i, sl in enumerate(slots):
        sl.idx = i
    return slots


def apply_cancellations(slots, cancels):
    for c in cancels:
        st = stand(c["stand"]) if c.get("stand") else None
        for sl in slots:
            if st and sl.stand != st:
                continue
            if "start" in c:
                hit = sl.start == parse_dt(c["start"])
            elif "time" in c:
                hit = sl.start.strftime("%H:%M") == c["time"]
                if "date" in c:  # date = the calendar date of the slot start
                    hit = hit and sl.start.date().isoformat() == c["date"]
            elif "from" in c:
                hit = parse_dt(c["from"]) <= sl.start and sl.end <= parse_dt(c["to"])
            else:
                raise SystemExit(f"Bad cancel entry: {c}")
            if hit:
                sl.cancelled = True


class Person:
    def __init__(self, pid, rules, extra=False):
        self.id = pid
        self.extra = extra
        self.unavailable = [(parse_dt(a), parse_dt(b)) for a, b in rules.get("unavailable", [])]
        self.only = {stand(s) for s in rules.get("stands", [])} or None
        self.never = {stand(s) for s in rules.get("not_stands", [])}
        self.night_stands = {stand(s) for s in rules.get("night_stands", [])} or None
        self.night_only = rules.get("night_only", False)
        self.no_night = rules.get("no_night", False)
        self.never_hours = [(int(a[:2]), int(b[:2])) for a, b in rules.get("never_hours", [])]
        self.max_shifts = rules.get("max_shifts")
        self.allow_same = rules.get("allow_same_stand", False)

    def can_take(self, sl):
        if self.only and sl.stand not in self.only:
            return False
        if sl.stand in self.never:
            return False
        if self.night_stands and sl.night is not None and sl.stand not in self.night_stands:
            return False
        if self.night_only and sl.night is None:
            return False
        if self.no_night and sl.night is not None:
            return False
        for a, b in self.never_hours:  # daily window, e.g. 05:00-07:00
            day = sl.start.replace(hour=0, minute=0)
            for d in (day - 24 * H, day):
                ws, we = d + a * H, d + (b if b > a else b + 24) * H
                if sl.start < we and sl.end > ws:
                    return False
        return all(sl.end <= a or sl.start >= b for a, b in self.unavailable)


def load(cfg):
    start, end = parse_dt(cfg["start"]), parse_dt(cfg["end"]) if "end" in cfg else None
    if end is None:
        end = start + dt.timedelta(days=cfg.get("days", 1))
    rules = {str(k): v for k, v in cfg.get("rules", {}).items()}
    ids = cfg.get("roster") or [str(i) for i in range(1, int(cfg["people"]) + 1)]
    people = [Person(str(i), rules.get(str(i), {})) for i in ids]
    for ex in cfg.get("extra", []):
        r = dict(ex)
        r.update(rules.get(str(ex["id"]), {}))
        people.append(Person(str(ex["id"]), r, extra=True))
    slots = build_slots(start, end)
    apply_cancellations(slots, cfg.get("cancel", []))
    return start, end, people, slots


def shortage_report(people, slots):
    by_night = defaultdict(list)
    for sl in slots:
        if sl.night is not None and not sl.cancelled:
            by_night[sl.night].append(sl)
    short = []
    for n, sls in sorted(by_night.items()):
        avail = [p for p in people if any(p.can_take(s) for s in sls)]
        if len(sls) > len(avail):
            short.append({
                "night": f"{n:%d.%m} 23:00 -> {n + dt.timedelta(days=1):%d.%m} 07:00",
                "slots": len(sls), "people_available": len(avail),
                "missing": len(sls) - len(avail),
            })
    return short


def solve_once(people, slots, cfg, rng, params):
    min_rest = cfg.get("min_rest_hours", 2) * H
    night_rest = cfg.get("night_rest_hours", 0) * H
    last_night = {}
    global_same = cfg.get("allow_same_stand", False)
    alpha, jitter = params
    last_end, last_stand = {}, {}
    for pid, ls in cfg.get("last_shift", {}).items():
        last_end[str(pid)] = parse_dt(ls["end"])
        last_stand[str(pid)] = stand(ls["stand"]) if ls.get("stand") else None
        last_night[str(pid)] = night_id(parse_dt(ls["end"]) - H) is not None
    hours = defaultdict(float)
    nights = defaultdict(int)
    count = defaultdict(int)
    used_night = defaultdict(set)
    assign = {}
    fixed = {}
    fixed_by_p = defaultdict(list)
    for f in cfg.get("fixed", []):
        st, s = stand(f["stand"]), parse_dt(f["start"])
        for sl in slots:
            if sl.stand == st and sl.start == s:
                fixed[sl.idx] = str(f["id"])
                fixed_by_p[str(f["id"])].append(sl)
    pmap = {p.id: p for p in people}

    def commit(pid, sl):
        assign[sl.idx] = pid
        last_end[pid], last_stand[pid] = sl.end, sl.stand
        last_night[pid] = sl.night is not None
        hours[pid] += sl.hours
        count[pid] += 1
        if sl.night is not None:
            nights[pid] += 1
            used_night[sl.night].add(pid)

    i = 0
    active = [s for s in slots if not s.cancelled]
    while i < len(active):
        j = i
        while j < len(active) and active[j].start == active[i].start:
            j += 1
        group = active[i:j]
        i = j
        for sl in group:
            if sl.idx in fixed:
                commit(fixed[sl.idx], sl)
        free = [sl for sl in group if sl.idx not in fixed]
        rng.shuffle(free)
        taken = {assign[s.idx] for s in group if s.idx in assign}
        for sl in free:
            best, best_score = None, None
            for p in people:
                pid = p.id
                if pid in taken or not p.can_take(sl):
                    continue
                if p.max_shifts is not None and count[pid] >= p.max_shifts:
                    continue
                le = last_end.get(pid)
                if le is not None and sl.start - le < min_rest:
                    continue
                if le is not None and last_night.get(pid) and sl.start - le < night_rest:
                    continue
                if sl.night is not None and pid in used_night[sl.night]:
                    continue
                if (not global_same and not p.allow_same
                        and last_stand.get(pid) == sl.stand):
                    continue
                bad = False
                for fs in fixed_by_p.get(pid, ()):
                    if fs.start >= sl.start and (fs.start - sl.end < min_rest
                                                 or (sl.night is not None and fs.night == sl.night)
                                                 or (fs.stand == sl.stand and not global_same and not p.allow_same)):
                        bad = True
                        break
                if bad:
                    continue
                idle = 1000.0 if le is None else (sl.start - le) / H
                score = idle - alpha * hours[pid] + rng.random() * jitter
                if sl.night is not None:
                    score -= alpha * 4 * nights[pid]
                if best_score is None or score > best_score:
                    best, best_score = pid, score
            if best is None:
                assign[sl.idx] = None
            else:
                commit(best, sl)
                taken.add(best)
    return assign


def evaluate(people, slots, assign, cfg):
    per = defaultdict(list)
    unfilled = 0
    for sl in slots:
        if sl.cancelled:
            continue
        pid = assign.get(sl.idx)
        if pid is None:
            unfilled += 1
        else:
            per[pid].append(sl)
    gaps_all, per_min = [], {}
    for pid, sls in per.items():
        sls.sort(key=lambda s: s.start)
        prev = None
        ls = cfg.get("last_shift", {}).get(pid)
        if ls:
            prev = parse_dt(ls["end"])
        g = []
        for s in sls:
            if prev is not None:
                g.append((s.start - prev) / H)
            prev = s.end
        gaps_all += g
        per_min[pid] = min(g) if g else None
    regular = [p.id for p in people if not p.extra]
    hrs = [sum(s.hours for s in per.get(pid, [])) for pid in regular]
    nts = [sum(1 for s in per.get(pid, []) if s.night is not None) for pid in regular]
    min_gap = min(gaps_all) if gaps_all else 99.0
    n_at_min = sum(1 for g in gaps_all if g == min_gap)
    score = (
        unfilled,
        -min_gap,
        n_at_min,
        (max(hrs) - min(hrs)) if hrs else 0,
        (max(nts) - min(nts)) if nts else 0,
        -sum(v for v in per_min.values() if v is not None),
    )
    return score, per, per_min


def solve(people, slots, cfg):
    rng = random.Random(cfg.get("seed"))
    budget = cfg.get("seconds", 8)
    max_iter = cfg.get("iterations", 100000)
    t0 = time.time()
    best = None
    it = 0
    while it < max_iter and (time.time() - t0) < budget:
        params = (rng.uniform(0, 0.6), rng.uniform(0.01, 3.0))
        a = solve_once(people, slots, cfg, rng, params)
        sc, per, per_min = evaluate(people, slots, a, cfg)
        if best is None or sc < best[0]:
            best = (sc, a, per, per_min)
        it += 1
    return best, it


# ---------------------------------------------------------------- output

def label(pid, names):
    if pid is None:
        return "???"
    n = names.get(pid)
    return f"{pid} - {n}" if n else pid


def write_xlsx(path, title, start, end, slots, assign, per, per_min, people, names):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    night_fill = PatternFill("solid", fgColor="D9E1F2")
    off_fill = PatternFill("solid", fgColor="E7E6E6")
    cancel_fill = PatternFill("solid", fgColor="F2F2F2")
    miss_fill = PatternFill("solid", fgColor="FFC7CE")

    # One sheet in the unit's layout (right to left):
    # יום | שעה | ש.ג | יוסטור | דרומית | פטרול | תצפית
    # Daytime 4h shifts go under תצפית (2 merged rows), night 2h shifts under פטרול.
    green = PatternFill("solid", fgColor="00FF00")
    grey_line = Side(style="medium", color="BFBFBF")
    ws = wb.create_sheet("רשימת שמירה")
    ws.sheet_view.rightToLeft = True
    heads = ["יום", "שעה", "ש.ג", "יוסטור", "דרומית", "פטרול", "תצפית"]
    col_of = {SG: 3, YUS: 4, DAR: 5}
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(size=12)
        cell.alignment = center
        cell.border = Border(bottom=grey_line)
    ws.row_dimensions[1].height = 28
    t0 = start.replace(minute=0)
    if (t0.hour - DAY_START_HOUR) % 2:
        t0 -= H
    nrows = 0
    t = t0
    while t < end:
        row = 2 + nrows
        ws.cell(row=row, column=2, value=f"{t:%H:%M}")
        for c in range(1, 8):
            x = ws.cell(row=row, column=c)
            x.alignment = center
            x.border = border
        ws.row_dimensions[row].height = 20
        nrows += 1
        t += 2 * H
    # day column, merged per 07:00-07:00 day
    day_letters = ["ב'", "ג'", "ד'", "ה'", "ו'", "ש'", "א'"]
    rows_by_day = defaultdict(list)
    for r in range(nrows):
        rows_by_day[day_id(t0 + 2 * r * H)].append(2 + r)
    for d, rows in rows_by_day.items():
        if len(rows) > 1:
            ws.merge_cells(start_row=rows[0], start_column=1, end_row=rows[-1], end_column=1)
        cell = ws.cell(row=rows[0], column=1, value=day_letters[d.weekday()])
        cell.font = Font(size=14)
        cell.alignment = center
    for sl in slots:
        if sl.stand == PAT:
            col = 6 if sl.night is not None else 7
        else:
            col = col_of[sl.stand]
        r0 = 2 + int((sl.start - t0) / H) // 2
        r1 = 2 + (int((sl.end - t0) / H) - 1) // 2
        if r1 > r0:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
        cell = ws.cell(row=r0, column=col)
        if sl.cancelled:
            cell.value, cell.fill = "מבוטל", cancel_fill
            cell.font = Font(italic=True, color="808080")
        else:
            pid = assign.get(sl.idx)
            cell.value = int(pid) if pid and pid.isdigit() and pid not in names else label(pid, names)
            cell.fill = miss_fill if pid is None else green
            cell.font = Font(size=12)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 12
    for c in "CDEFG":
        ws.column_dimensions[c].width = 18

    ws = wb.create_sheet("סיכום לפי מספר")
    ws.sheet_view.rightToLeft = True
    heads = ["מספר", "שם", "משמרות", "שעות", "משמרות לילה", "מנוחה מינ' (ש')", "פירוט"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = center
    for r, p in enumerate(people, 2):
        sls = per.get(p.id, [])
        detail = " | ".join(f"{s.start:%d.%m} {s.start:%H:%M}-{s.end:%H:%M} {s.stand}" for s in sls)
        vals = [int(p.id) if p.id.isdigit() else p.id, names.get(p.id, ""), len(sls),
                sum(s.hours for s in sls), sum(1 for s in sls if s.night is not None),
                per_min.get(p.id), detail]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(horizontal="center" if c < 7 else "right", vertical="center")
    for c, w in zip("ABCDEFG", [8, 16, 9, 8, 12, 14, 90]):
        ws.column_dimensions[c].width = w
    wb.save(path)


def write_whatsapp(path, title, slots, assign, names):
    lines = []
    days = sorted({day_id(s.start) for s in slots})
    for d in days:
        lines.append(f"*{title} — {HEB_DAYS[d.weekday()]} {d:%d.%m}*")
        for st in STANDS:
            sls = [s for s in slots if day_id(s.start) == d and s.stand == st]
            if not sls:
                continue
            lines.append("")
            lines.append(f"*{st}*")
            for s in sls:
                who = "מבוטל" if s.cancelled else label(assign.get(s.idx), names)
                lines.append(f"{s.start:%H:%M}-{s.end:%H:%M}  ←  {who}")
        lines.append("")
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines).strip() + "\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    ap.add_argument("--out", default=".")
    args = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)
    start, end, people, slots = load(cfg)
    names = {str(k): v for k, v in cfg.get("names", {}).items()}

    short = shortage_report(people, slots)
    if short and not cfg.get("ignore_shortage"):
        print(json.dumps({"status": "SHORTAGE", "nights": short}, ensure_ascii=False, indent=2))
        sys.exit(2)

    (score, assign, per, per_min), iters = solve(people, slots, cfg)
    title = cfg.get("title", "רשימת שמירה")
    os.makedirs(args.out, exist_ok=True)
    stem = cfg.get("file_name") or f"shmira_{start:%Y-%m-%d}"
    xlsx = os.path.join(args.out, stem + ".xlsx")
    txt = os.path.join(args.out, stem + "_whatsapp.txt")
    write_xlsx(xlsx, title, start, end, slots, assign, per, per_min, people, names)
    write_whatsapp(txt, title, slots, assign, names)

    state = {}
    for pid, sls in per.items():
        last = max(sls, key=lambda s: s.end)
        state[pid] = {"end": f"{last.end:%Y-%m-%d %H:%M}", "stand": last.stand}
    with open(os.path.join(args.out, stem + "_state.json"), "w", encoding="utf-8") as f:
        json.dump({"last_shift": state}, f, ensure_ascii=False, indent=2)

    regular = [p for p in people if not p.extra]
    hrs = [sum(s.hours for s in per.get(p.id, [])) for p in regular]
    unfilled = [f"{s.start:%d.%m %H:%M} {s.stand}" for s in slots
                if not s.cancelled and assign.get(s.idx) is None]
    report = {
        "status": "UNFILLED" if unfilled else "OK",
        "unfilled_slots": unfilled,
        "files": [xlsx, txt],
        "iterations": iters,
        "min_rest_hours": -score[1],
        "hours_per_person_min_max": [min(hrs), max(hrs)] if hrs else None,
        "night_shifts_spread": score[4],
        "people_with_no_shift": [p.id for p in people if p.id not in per],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    sys.exit(3 if unfilled else 0)


if __name__ == "__main__":
    main()
